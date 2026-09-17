"""K-에듀파인 개인수신그룹 일괄등록 지원.

에듀파인 [개인설정 > 개인수신그룹관리 > 일괄등록] 이 받는 엑셀을 만들고,
같은 화면의 [파일양식받기] 로 내려받은 파일에서 기관코드를 거꾸로 수확한다.

양식은 시트 '개인수신그룹관리', 1행 헤더, A~I 9개 열이다.
서버가 Apache POI 로 읽으므로 원본 양식을 열어 값만 채운다 (새로 만들지 않는다).
"""

import datetime
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:                                  # pragma: no cover
    openpyxl = None

from sotong_parser import lookup_org

SHEET_NAME = '개인수신그룹관리'
HEADERS = [
    '수신그룹명', '수신그룹기호', '수신그룹등록교육청', '수신기관구분',
    '수신기관코드', '수신기관명', '수신그룹유형', '수신사용자ID', '수신사용자명',
]
COL_GROUP_OFFICE = 2   # C 수신그룹등록교육청
COL_ORG_KIND = 3       # D 수신기관구분
COL_ORG_CODE = 4       # E 수신기관코드
COL_ORG_NAME = 5       # F 수신기관명
COL_USER_ID = 7        # H 수신사용자ID
COL_USER_NAME = 8      # I 수신사용자명

KIND_INTERNAL = 'NS'   # 대내조직 (MG:LDAP, MP:문서24조직, 99:사용자)
TYPE_DOCUMENT = '02'   # 문서 (01:메모, 04:메일, 05:공람, 06:일정)

# 에듀파인 기관코드는 영문자 1글자 + 숫자 9자리 (예: M100000795, 충북은 M10 으로 시작)
CODE_RE = re.compile(r'^[A-Z]\d{9}$')

# 양식의 안내용 예시 행에 들어 있는 표현. 수확할 때 걸러낸다.
PLACEHOLDER_HINTS = ('(수신', '수신그룹명', '수신기관코드', '수신기관명')


def _resource_path(name: str) -> str:
    base = getattr(sys, '_MEIPASS', os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, name)


CODES_FILE = _resource_path('org_codes.json')
TEMPLATE_FILE = _resource_path(os.path.join('assets', '수신그룹_양식.xlsx'))


def _require_openpyxl():
    if openpyxl is None:
        raise RuntimeError('openpyxl 이 설치되어 있지 않습니다. pip install openpyxl')


# ── 기관코드 사전 ────────────────────────────

def empty_codes() -> dict:
    # '기관' 은 {에듀파인 수신기관명(전체경로): 기관코드}.
    # 짧은 이름을 키로 쓰면 '행정지원과' 처럼 여러 기관에 같은 부서명이 있을 때
    # 뒤엣것이 앞엣것을 덮어써 조용히 틀린 코드가 박힌다. 전체경로는 유일하다.
    return {'등록교육청코드': '', '수집일': '', '기관': {}}


def load_codes(path: str = None) -> dict:
    path = path or CODES_FILE
    if not os.path.exists(path):
        return empty_codes()
    with open(path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    base = empty_codes()
    base.update(data)
    return base


def save_codes(codes: dict, path: str = None) -> str:
    path = path or CODES_FILE
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(codes, f, ensure_ascii=False, indent=2, sort_keys=True)
        f.write('\n')
    return path


def short_name(full_name: str) -> str:
    """에듀파인 수신기관명 → 사람이 부르는 이름.

    '충청북도진천교육지원청 학성초등학교' → '학성초등학교'
    '충청북도교육청 유초등교육과'          → '유초등교육과'
    org_db 에 있는 이름이면 정식명으로 맞춘다.
    """
    tail = (full_name or '').strip().split(' ')[-1].strip()
    if not tail:
        return ''
    return lookup_org(tail) or tail


def read_group_workbook(path: str) -> dict:
    """에듀파인에서 받은 양식/등록 파일 → {'등록교육청코드', '사용자ID', '사용자명', '기관': {...}}

    안내용 예시 행과 빈 행은 건너뛴다. 코드 형식이 맞는 행만 받는다.
    """
    _require_openpyxl()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]

    out = {'등록교육청코드': '', '사용자ID': '', '사용자명': '', '기관': {}}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or len(row) <= COL_ORG_NAME:
            continue
        code = str(row[COL_ORG_CODE] or '').strip()
        full = str(row[COL_ORG_NAME] or '').strip()
        if not CODE_RE.match(code) or not full:
            continue
        if any(h in full for h in PLACEHOLDER_HINTS):
            continue

        out['기관'][full] = code

        office = str(row[COL_GROUP_OFFICE] or '').strip()
        if CODE_RE.match(office) and not out['등록교육청코드']:
            out['등록교육청코드'] = office
        if not out['사용자ID'] and row[COL_USER_ID]:
            out['사용자ID'] = str(row[COL_USER_ID]).strip()
        if not out['사용자명'] and row[COL_USER_NAME]:
            out['사용자명'] = str(row[COL_USER_NAME]).strip()
    return out


def merge_codes(codes: dict, harvested: dict) -> tuple:
    """수확 결과를 사전에 합친다. (갱신된 사전, 새로 추가된 수, 코드가 바뀐 기관 목록)"""
    codes = dict(codes)
    codes.setdefault('기관', {})
    orgs = dict(codes['기관'])

    added, changed = 0, []
    for full, code in harvested.get('기관', {}).items():
        old = orgs.get(full)
        if old is None:
            added += 1
        elif old != code:
            changed.append((full, old, code))
        orgs[full] = code

    codes['기관'] = orgs
    if harvested.get('등록교육청코드'):
        codes['등록교육청코드'] = harvested['등록교육청코드']
    codes['수집일'] = datetime.date.today().isoformat()
    return codes, added, changed


def index_by_short_name(codes: dict) -> dict:
    """{짧은 이름: [전체경로, ...]}. 같은 짧은 이름이 여러 기관에 있으면 목록이 길어진다."""
    index = {}
    for full in (codes or {}).get('기관', {}):
        for key in {short_name(full), full}:
            if key:
                index.setdefault(key, [])
                if full not in index[key]:
                    index[key].append(full)
    return index


def lookup_code(codes: dict, name: str, index: dict = None):
    """이름 → ({'code','fullname'} 또는 None, 동명 후보 전체경로 목록).

    후보가 둘 이상이면 아무것도 고르지 않고 후보만 돌려준다.
    임의로 하나를 고르면 공문이 엉뚱한 부서로 간다.
    """
    orgs = (codes or {}).get('기관', {})
    if not orgs or not name:
        return None, []

    if name in orgs:                       # 전체경로로 직접 지목
        return {'code': orgs[name], 'fullname': name}, []

    index = index if index is not None else index_by_short_name(codes)
    matches = index.get(name) or index.get(lookup_org(name) or '') or []
    if len(matches) == 1:
        full = matches[0]
        return {'code': orgs[full], 'fullname': full}, []
    return None, list(matches)


# ── 일괄등록 엑셀 만들기 ──────────────────────

def safe_filename(text: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]', '', (text or '').strip())
    return cleaned or '수신그룹'


def build_workbook(rows, meta, out_path, template_path=None):
    """수신그룹 일괄등록 엑셀을 만든다.

    rows: [{'name': 정식명, 'code': 기관코드, 'fullname': 에듀파인 수신기관명}, ...]
    meta: {'그룹명', '그룹기호', '등록교육청코드', '사용자ID', '사용자명',
           '기관구분'(선택), '그룹유형'(선택)}
    """
    _require_openpyxl()
    if not rows:
        raise ValueError('등록할 기관이 없습니다.')
    for key in ('그룹명', '등록교육청코드', '사용자ID', '사용자명'):
        if not str(meta.get(key, '')).strip():
            raise ValueError(f'{key} 이(가) 비어 있습니다.')

    template_path = template_path or TEMPLATE_FILE
    wb = openpyxl.load_workbook(template_path)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]

    # 양식에 안내용 예시 행이 남아 있을 수 있으므로 헤더만 남기고 비운다
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)

    kind = meta.get('기관구분') or KIND_INTERNAL
    gtype = meta.get('그룹유형') or TYPE_DOCUMENT
    for row in rows:
        ws.append([
            meta['그룹명'],
            meta.get('그룹기호') or ' ',
            meta['등록교육청코드'],
            kind,
            row['code'],
            row.get('fullname') or row['name'],
            gtype,
            meta['사용자ID'],
            meta['사용자명'],
        ])

    wb.save(out_path)
    return out_path


def default_output_name(group_name: str, when=None) -> str:
    when = when or datetime.date.today()
    return f'수신그룹_{safe_filename(group_name)}_{when:%Y%m%d}.xlsx'


def split_by_code(parsed_rows, codes):
    """parse_orgs 결과 → (코드 있는 행, 코드 없는 행).

    확정되지 않은 행(fuzzy/none)은 애초에 여기 들어오면 안 된다.
    코드가 없거나 동명 기관이 여럿이면 조용히 빠지지 않고 사유와 함께 돌려준다.
    """
    index = index_by_short_name(codes)
    ready, missing = [], []
    for row in parsed_rows:
        name = row.get('name')
        if not name:
            missing.append(dict(row, reason='이름 미확정'))
            continue
        entry, ambiguous = lookup_code(codes, name, index)
        if entry:
            ready.append({'name': name, 'code': entry['code'],
                          'fullname': entry['fullname']})
        elif ambiguous:
            missing.append(dict(row, reason='동명 기관 여럿', candidates=ambiguous))
        else:
            missing.append(dict(row, reason='코드 없음'))
    return ready, missing
