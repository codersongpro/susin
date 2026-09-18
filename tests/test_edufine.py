import os
import tempfile
import unittest
from unittest import mock

import edufine

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLE = os.path.join(HERE, 'data', 'sample_수신그룹.xlsx')
TEMPLATE = os.path.join(HERE, 'assets', '수신그룹_양식.xlsx')


def codes_from_sample():
    harvested = edufine.read_group_workbook(SAMPLE)
    codes, _, _ = edufine.merge_codes(edufine.empty_codes(), harvested)
    return codes


class HarvestTest(unittest.TestCase):
    def test_reads_real_codes_and_owner(self):
        h = edufine.read_group_workbook(SAMPLE)
        self.assertEqual(h['등록교육청코드'], 'M100000098')
        self.assertEqual(h['사용자ID'], 'dungst')
        self.assertEqual(h['기관']['충청북도진천교육지원청 학성초등학교'], 'M100000795')

    def test_blank_template_yields_nothing(self):
        # 안내용 예시 행이 코드로 섞여 들어가면 안 된다
        self.assertEqual(edufine.read_group_workbook(TEMPLATE)['기관'], {})

    def test_merge_reports_added_and_changed(self):
        codes, added, changed = edufine.merge_codes(
            edufine.empty_codes(), edufine.read_group_workbook(SAMPLE))
        self.assertEqual(added, 4)
        self.assertEqual(changed, [])

        moved = {'기관': {'충청북도진천교육지원청 학성초등학교': 'M100009999'}}
        _, added2, changed2 = edufine.merge_codes(codes, moved)
        self.assertEqual(added2, 0)
        self.assertEqual(changed2,
                         [('충청북도진천교육지원청 학성초등학교', 'M100000795', 'M100009999')])


class CodePersistenceTest(unittest.TestCase):
    def test_user_file_overrides_bundled_file_and_is_saved_atomically(self):
        with tempfile.TemporaryDirectory() as tmp:
            bundled = os.path.join(tmp, 'bundled.json')
            user = os.path.join(tmp, 'user', 'org_codes.json')
            base = {'등록교육청코드': '', '수집일': '',
                    '기관': {'기본 기관': 'M100000001'}}
            updated = {'등록교육청코드': '', '수집일': '2026-09-18',
                       '기관': {'갱신 기관': 'M100000002'}}
            edufine.save_codes(base, bundled)

            with mock.patch.object(edufine, 'BUNDLED_CODES_FILE', bundled), \
                    mock.patch.object(edufine, 'USER_CODES_FILE', user):
                self.assertEqual(edufine.load_codes()['기관'], base['기관'])
                saved = edufine.save_codes(updated)
                self.assertEqual(saved, user)
                self.assertTrue(os.path.exists(user))
                self.assertFalse(os.path.exists(user + '.tmp'))
                self.assertEqual(edufine.load_codes()['기관'], updated['기관'])


class LookupTest(unittest.TestCase):
    def setUp(self):
        self.codes = codes_from_sample()

    def test_lookup_by_alias_and_full_name(self):
        for q in ('학성초', '학성초등학교', '충청북도진천교육지원청 학성초등학교'):
            entry, ambiguous = edufine.lookup_code(self.codes, q)
            self.assertEqual(entry['code'], 'M100000795', q)
            self.assertEqual(ambiguous, [])

    def test_unknown_returns_nothing(self):
        self.assertEqual(edufine.lookup_code(self.codes, '없는기관'), (None, []))

    def test_duplicate_short_name_is_never_guessed(self):
        # 부서명은 기관마다 겹친다. 하나를 임의로 고르면 공문이 엉뚱한 곳으로 간다.
        codes = {'기관': {
            '충청북도청주교육지원청 행정지원과': 'M100000001',
            '충청북도충주교육지원청 행정지원과': 'M100000002',
        }}
        entry, ambiguous = edufine.lookup_code(codes, '행정지원과')
        self.assertIsNone(entry)
        self.assertEqual(len(ambiguous), 2)

        # 전체경로로 지목하면 정확히 찾는다
        entry, _ = edufine.lookup_code(codes, '충청북도충주교육지원청 행정지원과')
        self.assertEqual(entry['code'], 'M100000002')


class SplitTest(unittest.TestCase):
    def test_missing_rows_are_reported_with_reason(self):
        codes = codes_from_sample()
        parsed = [
            {'name': '학성초등학교', 'grade': 'exact'},
            {'name': '백곡초등학교', 'grade': 'exact'},   # 코드 사전에 없음
        ]
        ready, missing = edufine.split_by_code(parsed, codes)
        self.assertEqual([r['name'] for r in ready], ['학성초등학교'])
        self.assertEqual([r['name'] for r in missing], ['백곡초등학교'])
        self.assertEqual(missing[0]['reason'], '코드 없음')


class BuildTest(unittest.TestCase):
    meta = {'그룹명': '진천 시범', '그룹기호': '', '등록교육청코드': 'M100000098',
            '사용자ID': 'dungst', '사용자명': '송동석'}

    def test_round_trip(self):
        codes = codes_from_sample()
        rows, _ = edufine.split_by_code(
            [{'name': '학성초등학교'}, {'name': '한천초등학교'}], codes)
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, 'out.xlsx')
            edufine.build_workbook(rows, self.meta, out)

            import openpyxl
            wb = openpyxl.load_workbook(out)
            self.assertIn(edufine.SHEET_NAME, wb.sheetnames)
            ws = wb[edufine.SHEET_NAME]
            self.assertEqual([c.value for c in ws[1]], edufine.HEADERS)
            self.assertEqual(ws.max_row, 3)

            second = [c.value for c in ws[2]]
            self.assertEqual(second[0], '진천 시범')
            self.assertEqual(second[1], ' ')          # 기호가 비면 공백 한 칸
            self.assertEqual(second[2], 'M100000098')
            self.assertEqual(second[3], edufine.KIND_INTERNAL)
            self.assertEqual(second[4], 'M100000795')
            self.assertEqual(second[5], '충청북도진천교육지원청 학성초등학교')
            self.assertEqual(second[6], edufine.TYPE_DOCUMENT)
            self.assertEqual(second[7:], ['dungst', '송동석'])

    def test_generated_file_harvests_back(self):
        # 만든 파일을 다시 읽어 같은 코드가 나와야 한다
        codes = codes_from_sample()
        rows, _ = edufine.split_by_code([{'name': '학성초등학교'}], codes)
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, 'out.xlsx')
            edufine.build_workbook(rows, self.meta, out)
            back = edufine.read_group_workbook(out)
            self.assertEqual(back['기관']['충청북도진천교육지원청 학성초등학교'],
                             'M100000795')

    def test_refuses_empty_or_incomplete(self):
        with self.assertRaises(ValueError):
            edufine.build_workbook([], self.meta, 'x.xlsx')
        with self.assertRaises(ValueError):
            edufine.build_workbook(
                [{'name': 'a', 'code': 'M100000001'}],
                dict(self.meta, 사용자ID=''), 'x.xlsx')

    def test_default_output_name(self):
        import datetime
        self.assertEqual(
            edufine.default_output_name('진천/시범', datetime.date(2026, 9, 17)),
            '수신그룹_진천시범_20260917.xlsx')


if __name__ == '__main__':
    unittest.main()


class ResolveOrgTest(unittest.TestCase):
    """부서까지 지목할 수 있는지. 학교는 이름 하나로 끝나지만 부서는 아니다."""

    def setUp(self):
        self.codes = edufine.load_codes()
        self.index = edufine.index_by_short_name(self.codes)

    def resolve(self, text):
        return edufine.resolve_org(self.codes, text, self.index)

    def test_school_by_alias(self):
        full, amb = self.resolve('학성초')
        self.assertEqual(full, '충청북도진천교육지원청 학성초등학교')
        self.assertEqual(amb, [])

    def test_unique_department_by_name_alone(self):
        full, _ = self.resolve('정책기획과')
        self.assertEqual(full, '충청북도교육청 정책기획과')

    def test_shared_department_needs_its_parent(self):
        # '행정과'는 11곳에 있다. 혼자서는 절대 확정되면 안 된다.
        full, amb = self.resolve('행정과')
        self.assertIsNone(full)
        self.assertGreater(len(amb), 1)

    def test_parent_plus_department_narrows_to_one(self):
        self.assertEqual(self.resolve('청주교육지원청 행정과')[0],
                         '충청북도청주교육지원청 행정과')
        self.assertEqual(self.resolve('충주교육지원청 행정과')[0],
                         '충청북도충주교육지원청 행정과')

    def test_similar_parents_are_never_confused(self):
        # 청주와 충주는 한 글자 차이다. 편집거리로 붙이면 공문이 옆 시로 간다.
        for parent in ('청주', '충주', '제천', '단양', '보은', '옥천', '영동', '음성', '진천'):
            full, _ = self.resolve(f'{parent}교육지원청 행정과')
            self.assertIsNotNone(full, parent)
            self.assertIn(parent, full, parent)

    def test_nested_department(self):
        self.assertEqual(
            self.resolve('단재교육연수원 교육연수부')[0],
            '충청북도교육청 충청북도단재교육연수원 교육연수부')

    def test_full_path_is_taken_as_is(self):
        for full in ('충청북도청주교육지원청 행정과', '충청북도교육청 유초등교육과'):
            self.assertEqual(self.resolve(full), (full, []))

    def test_unknown_yields_nothing(self):
        self.assertEqual(self.resolve('있을리없는부서'), (None, []))

    def test_every_resolved_name_has_a_code(self):
        for text in ('학성초', '정책기획과', '청주교육지원청 행정과',
                     '단재교육연수원 교육연수부', '충북외고'):
            full, _ = self.resolve(text)
            entry, _ = edufine.lookup_code(self.codes, full, self.index)
            self.assertIsNotNone(entry, text)


class DisplayNameTest(unittest.TestCase):
    def setUp(self):
        self.codes = edufine.load_codes()

    def test_unique_short_name_is_shown_short(self):
        self.assertEqual(
            edufine.display_name(self.codes, '충청북도진천교육지원청 학성초등학교'),
            '학성초등학교')

    def test_shared_short_name_shows_full_path(self):
        full = '충청북도청주교육지원청 행정과'
        self.assertEqual(edufine.display_name(self.codes, full), full)


class SearchOrgsTest(unittest.TestCase):
    def setUp(self):
        self.codes = edufine.load_codes()

    def test_all_pieces_must_match(self):
        hits = edufine.search_orgs(self.codes, '청주 초등학교')
        self.assertTrue(hits)
        for h in hits:
            self.assertIn('청주', h)
            self.assertIn('초등학교', h)

    def test_department_search(self):
        hits = edufine.search_orgs(self.codes, '행정과')
        self.assertGreater(len(hits), 1)
        self.assertTrue(all('행정과' in h for h in hits))

    def test_empty_query_lists_everything(self):
        self.assertEqual(len(edufine.search_orgs(self.codes, '', limit=10000)),
                         len(self.codes['기관']))


class RegisteringOfficesTest(unittest.TestCase):
    def setUp(self):
        self.codes = edufine.load_codes()

    def test_lists_province_and_district_offices(self):
        offices = dict(edufine.registering_offices(self.codes))
        self.assertEqual(offices.get('충청북도교육청'), 'M100000001')
        self.assertEqual(offices.get('충청북도진천교육지원청'), 'M100000098')
        self.assertEqual(len(offices), 11)      # 도교육청 1 + 교육지원청 10

    def test_excludes_departments_and_schools(self):
        names = [n for n, _ in edufine.registering_offices(self.codes)]
        self.assertTrue(all(' ' not in n for n in names))
        self.assertNotIn('충청북도진천교육지원청 학성초등학교', names)

    def test_empty_dictionary_is_safe(self):
        self.assertEqual(edufine.registering_offices(edufine.empty_codes()), [])


class CategoryTest(unittest.TestCase):
    """찾아보기 분류 필터."""

    def setUp(self):
        self.codes = edufine.load_codes()

    def test_categorise(self):
        cases = {
            '충청북도진천교육지원청 학성초등학교': '초등학교',
            '충청북도청주교육지원청 청주중학교': '중학교',
            '충청북도교육청 충북외국어고등학교': '고등학교',
            '충청북도청주교육지원청 행정과': '부서·기관',
            '충청북도청주교육지원청': '교육지원청',
        }
        for full, want in cases.items():
            self.assertEqual(edufine.categorise(full), want, full)

    def test_counts_cover_everything(self):
        counts = edufine.category_counts(self.codes)
        self.assertEqual(sum(counts.values()), len(self.codes['기관']))

    def test_filtered_search(self):
        hits = edufine.search_orgs(self.codes, '청주', category='초등학교')
        self.assertTrue(hits)
        for h in hits:
            self.assertIn('청주', h)
            self.assertEqual(edufine.categorise(h), '초등학교')

    def test_category_alone_lists_them_all(self):
        offices = edufine.search_orgs(self.codes, '', category='교육지원청')
        self.assertEqual(len(offices), 10)
